"""Export the second-round review and verify the deliverable."""
from pathlib import Path
import sys,json,shutil
import pandas as pd,numpy as np
ROOT=Path.cwd();sys.path.insert(0,str(ROOT))
from backtest.engine import run_strategy
from backtest.run_manifest import artifact_record,write_manifest
RUN=Path(__file__).resolve().parents[1];O=RUN/'outputs';I=RUN/'inputs/data'
random=pd.read_csv(O/'random_calendar_summary.csv');counts=pd.read_csv(O/'random_flip_counts.csv')
events=pd.read_csv(O/'event_decomposition.csv');es=pd.read_csv(O/'event_decomposition_summary.csv')
focus=pd.read_csv(O/'focus_catalog.csv');review=pd.read_csv(O/'rejection_review.csv')
same=pd.read_csv(O/'same_mapping_comparisons.csv');old=pd.read_csv(I/'full_sensitivity.csv')
v=json.loads((O/'verification.json').read_text())
key='grid_ew_L20_zw250_sm0__carry__ew_sym'
k=es[es.candidate.eq(key)&es.lag.eq(1)].iloc[0]
original=old[old.candidate.eq(key)&old.lag.eq(1)&old['mask'].eq('all_10')].iloc[0]
first_share=k.first_day_log_gap/original.total_relative_log
later_share=k.later_days_log_gap/original.total_relative_log
remaining_share=original.nonevent_relative_log/original.total_relative_log
assert np.isclose(first_share+later_share+remaining_share,1)
same_lf=same[same.reference.eq('ew_lf')&same.lag.eq(1)&same['mask'].eq('all_10')]
assert len(same_lf)==24 and (same_lf.neutral_delta_sharpe<0).all()

# Exact daily split: candidate LF - EW symmetric = (candidate LF - EW LF) - EW short.
positions=pd.read_csv(I/'candidate_positions.csv',index_col=0,parse_dates=True)
refs=pd.read_csv(I/'reference_positions.csv',index_col=0,parse_dates=True)
market=pd.read_csv(I/'market.csv',index_col=0,parse_dates=True)
masks=pd.read_csv(I/'event_masks.csv',index_col=0,parse_dates=True)
split=[];maxsplit=0
for r in focus[focus.mapping.eq('lf')].itertuples():
    p=positions[r.name].dropna();idx=p.index
    for lag in [1,2]:
        def ret(q):
            q=q.reindex(idx)
            if lag==2:q=q.shift(1).fillna(0)
            return run_strategy(q,market.underlying,carry=market['carry']).ret
        cand=ret(p);long=ret(refs.ew_lf);short=ret(refs.ew_short);sym=ret(refs.ew_sym)
        err=float((sym-long-short).abs().max());maxsplit=max(maxsplit,err);assert err<1e-12
        for segment,sel in [('full',np.ones(len(idx),bool)),('events',masks.all_10.reindex(idx).to_numpy()),('non_events',~masks.all_10.reindex(idx).to_numpy())]:
            longgap=(cand-long).loc[sel].sum();missing_short=-short.loc[sel].sum();actual=(cand-sym).loc[sel].sum()
            assert np.isclose(longgap+missing_short,actual)
            split.append(dict(candidate=r.candidate,lag=lag,segment=segment,long_signal_simple_gap=longgap,absent_short_simple_gap=missing_short,total_simple_gap=actual))
pd.DataFrame(split).to_csv(O/'long_signal_vs_short_leg.csv',index=False)

# Preserve total-family diagnostic distributions without labeling them calibrated p-values.
countsummary=[]
for (scope,lag),g in counts.groupby(['comparison_scope','lag']):
    eligible=random[random.comparison_scope.eq(scope)&random.lag.eq(lag)]
    ids=set(eligible.candidate)
    z=old[old.candidate.isin(ids)&old.lag.eq(lag)&old['mask'].eq('all_10')]
    actual=int(z.rank_flip.sum())
    countsummary.append(dict(comparison_scope=scope,lag=lag,n=len(eligible),actual_event_sharpe_flips=actual,
       random_flip_median=g.sharpe_flips.median(),random_flip_q05=g.sharpe_flips.quantile(.05),random_flip_q95=g.sharpe_flips.quantile(.95),
       random_at_least_actual_fraction=(g.sharpe_flips>=actual).mean()))
countsummary=pd.DataFrame(countsummary);countsummary.to_csv(O/'random_count_summary.csv',index=False)

# Standalone scientific figure: the neutral Sharpe gap under random calendar removal.
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
trial=pd.read_csv(O/'focus_random_trials.csv');t=trial[trial.candidate.eq(key)]
fig,axes=plt.subplots(1,2,figsize=(11,4),sharex=True,sharey=True,layout='constrained')
for lag,ax in zip([1,2],axes):
    values=t[t.lag.eq(lag)].neutral_delta_sharpe
    act=random[random.candidate.eq(key)&random.lag.eq(lag)].iloc[0]
    ax.hist(values,bins=np.linspace(-.24,.16,61),color='#527ca8',alpha=.85,density=True)
    ax.axvline(0,color='#777777',linewidth=1,label='Equal Sharpe')
    ax.axvline(act.actual_neutral_delta_sharpe,color='#d26936',linewidth=2,linestyle='--',label='Actual event windows')
    ax.set_title('T close execution' if lag==1 else 'T+1 close execution')
    ax.set_xlabel('Candidate minus incumbent Sharpe after removal')
    ax.text(.97,.95,f'Random flip rate: {act.random_sharpe_flip_fraction:.1%}',transform=ax.transAxes,ha='right',va='top')
axes[0].set_ylabel('Density');axes[0].legend(loc='upper left',fontsize=8)
fig.suptitle('EW 20/250/0: 2,000 same-year random 70-day removals')
fig.savefig(O/'random_calendar_diagnostic.png',dpi=180);fig.savefig(O/'random_calendar_diagnostic.svg');plt.close(fig)

with pd.ExcelWriter(O/'事件依赖第二轮_候选复核.xlsx',engine='openpyxl') as writer:
    for name,df in [('复核结论',review),('同映射比较',same),('事件方向与响应',events),('事件归因汇总',es),('空头腿与信号拆分',pd.DataFrame(split)),('随机对照重点候选',random[random.candidate.isin(focus.candidate)]),('随机对照全部',random),('随机反转数分布',countsummary),('逐年表现',pd.read_csv(O/'yearly_results.csv'))]:
        df.to_excel(writer,sheet_name=name,index=False);ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(15,len(str(col[0].value))+3))
detail=events[events.candidate.eq(key)&events.lag.eq(1)]
eventtable='|事件|候选/现役事前仓位|首日相对对数收益差|后9日差|事后仓位说明|\n|---|---|---:|---:|---|\n'
for r in detail.itertuples():
    desc=''
    if r.event=='tariff_20190802':desc='现役7/25已转空，8/7转多；候选6/12以来持多'
    elif r.event=='tariff_20190506':desc='初始都空，候选5/13开始变化'
    elif r.event=='geneva_20250512':desc='初始都多，候选5/15开始变化'
    eventtable+=f'|{r.event}|{r.candidate_initial:+.0f} / {r.reference_initial:+.0f}|{r.first_day_log_gap*100:+.3f}|{r.later_days_log_gap*100:+.3f}|{desc}|\n'

report=f'''# 事件依赖第二轮：哪些值得重审？

分析日期2026-09-15，价格截至2026-09-11。全部沿用r1冻结数据，不改现役、不重扫参数。本轮复核25个主窗口夏普反转规格，另加20日窗口中出现双指标反转的一组既有权重，共26个规格。未补齐r1的8项数值覆盖缺口。

## 核心结论

**上一轮的25个夏普反转不能直接成为25个“被事件错杀”的信号。** 24个涉及不同持仓结构；放回同多头基准后，事件中性夏普全部仍较差。剩余同类EW 20/250/0的轻微反转，在随机移除行情时也很常见，而且对执行时间敏感。

据此，本轮没有新增“事件剔除后应替换现役”的候选。保留两项敏感性观察：EW 20/250/0和原四对权重(2/7,1/7,2/7,2/7)。这不是证明所有未采用信号无效，也不是宣布现役参数最优。

## 1. 先把持仓结构放在同一把尺子上

上一轮25个规格中，24个是历史多头/空仓候选相对当前EW多空对称的反转。它们相对EW多头版本：

- 原全期夏普全部较低；
- 同样去掉70个事件窗口交易日后，夏普仍全部较低，差距为{abs(same_lf.neutral_delta_sharpe.max()):.3f}—{abs(same_lf.neutral_delta_sharpe.min()):.3f}；
- 因而不能用“相对多空现役反转”证明多头信号被事件错杀。

例如Hamilton多头剔事件后相对EW对称夏普+0.0357，但相对EW多头是−0.2156；单1000配对分别为+0.0422和−0.2091。这些是不同基准，不是数字矛盾。

收益上另作逐日精确分解：候选多头−EW对称 =（候选多头−EW多头）−EW空头腿。持仓、carry和交易成本均纳入。结果存入long_signal_vs_short_leg.csv。该拆分是简单日收益的加总，不能与复利对数收益混加。累计收益必须更高并非低风险策略有价值的必要条件；本轮排查的关键是同映射表现与原否决门槛。

## 2. 81.4%事件依赖，并不等于81.4%没猜中新闻

重点EW 20/250/0的原相对对数收益差为−0.121642。按发生时间拆开：

|差距发生在哪里|占全部相对落后的比例|
|---|---:|
|七项冲击首日|{first_share:.1%}|
|七项各自后9个交易日|{later_share:.1%}|
|事件窗口之外|{remaining_share:.1%}|

首日+后9日等于上一轮的81.4%。首日是A股反应锚日，包含全天收益，并非单独隔夜新闻跳空；后续也仍可包含事件影响。这里不能把“后9日”都解释成可预测收益。

{eventtable}

表中收益差单位为相对对数收益×100，负号表示候选落后，非简单百分比回报差。事前仓位的“生效日期”来自完整逐日账本，例如现役2019-07-25转空对应此前收盘信号。

2019年8月窗口最关键：现役冲击前已持空，随后8/7转多；候选一直持多。收益差同时包含初始方向与后续调整。按10日窗口简单收益差拆解，七事件合计：

|加总贡献（候选−现役）|简单收益百分点|
|---|---:|
|事前仓位差若维持不动产生的市场收益差|{k.initial_direction_simple_gap*100:+.3f}|
|双方事后相对事前仓位的变化贡献|{k.response_simple_gap*100:+.3f}|
|carry差|{k.carry_simple_gap*100:+.3f}|
|成本差|{k.cost_simple_gap*100:+.3f}|
|净差合计|{k.net_simple_gap*100:+.3f}|

这与上面的首日/后续日是两种切法，不能相加。冻结仓位贡献包含后续延续行情，不能证明是“运气”；响应贡献也不能证明系统识别了政策文本。双方法律/新闻事前知识不在数据里。

9.24和2025年4月窗口两者事前方向与窗口收益相同，均不是这个规格相对落后的来源。

## 3. 随机移除行情：这个排名反转是否特殊？

在相同年份随机取七段10日行情，年内段数保持2018两段、2019两段、2024一段、2025两段。同次不重叠、共70日，不跨年；每次所有候选采用相同日期，可覆盖真实事件。固定种子20260915，2000次。只对覆盖2018—2025完整年份的696组当前EW对称比较执行。

EW 20/250/0结果：

|指标|T收盘执行|次日收盘执行|
|---|---:|---:|
|原夏普差|−0.0158|−0.1533|
|真实事件窗口置零后夏普差|+0.0163|−0.1001|
|随机窗口置零后夏普反转比例|52.8%|0/2000|
|随机窗口置零后夏普与累计收益同时反转比例|27.3%|0/2000|
|随机窗口中相对收益损失不小于真实事件窗口的比例|34.25%|10.30%|

因此，在T收盘口径，真实事件造成的小幅夏普反转不是稀有现象；换成次日收盘执行，真实窗口和本次2000个随机窗口都没有反转。0/2000不等于理论概率为0。

![同年随机移除行情的夏普差分布](random_calendar_diagnostic.png)

不能把上述结论推广成“所有事件效应都不特殊”：在319个跨映射比较中，真实窗口有24个夏普反转，随机窗口的反转数中位数为4个；一些候选的真实事件损失也明显集中。这说明事件期的不同仓位结构确实重要，但这些候选仍未超过同多头基准。完整696组随机分布保留，未只展示有利案例。

上述比例是**事后随机日历诊断，不是显著性p值**：未校正候选挑选、事件清单选择或历史复用；事件发生日期也不能被视为随机分配。它只检验“同样移除70个交易日，排名变化是否常见”。

## 4. 原否决理由和重审清单

|对象|原否决/未采用的含义|本轮处理|
|---|---|---|
|EW 20/250/0|原网格未选参数；早期20/40/5选择轨迹不完整，不能虚构逐参数否决记录|保留为排名敏感、执行敏感的观察项；不据此升格|
|权重(2/7,1/7,2/7,2/7)|原权重族未通过选择校正；该行是上一轮20日窗口敏感对象|10日主窗口不反转；随机70日夏普反转39.45%，次日收盘仅0.65%；仍属脆弱排序|
|配对集合B/C、红利伙伴|原基准为long-flat；独立IC闸未过，收益要求worst(train,val)至少+0.15且换手/集中度不恶化|同多头基准未领先；不能将原2021—2023收益门槛失败归咎2024/2025事件|
|分批建仓|原为long-flat探索；验证期优势集中、收益回撤和进出机制均参与判断|四个本次反转规格在同多头基准仍较差；不是一次总收益排序即可推翻|
|动量、阈值、B2|原门槛涉及跨期、独立信息或选择校正；本轮反转行不一定是原代表|未发现因事件剔除而推翻原门槛的证据|
|raw/DEMA/Hamilton固定生成器|原预筛描述性，后续固定回放未形成替换依据|不冒称原正式STOP；继续作为既有对照|

原“被否决研究”和单一网格行并非同一单位。完整26行的原研究关联、同映射差距和复核说明见rejection_review.csv。此次没有重新执行原IC检验或全套选择校正，也没有给历史裁决补发统计证明。

## 5. 新维度怎样使用

事件依赖应当是一组诊断列，而不是一个越低越好的总分：

1. 同映射下事件前后排名，避免混入多空结构差异；
2. 事件相对损失占比，同时列原差距，避免接近零的分母放大比例；
3. 冲击首日、后续持仓响应和carry/成本分开；
4. 与同年随机窗口比较，识别一般性的样本敏感；
5. 5/10/20日窗口及两种执行时点的一致性；
6. 回到原训练/验证门槛，并等待真正新增的后续数据。

当前适合形成“重审优先级”，不足以挑出一个可替代现役的新赢家。不能从表现依赖事件推出纯运气，也不能从事件敏感性不强推出未来有效。

## 验证与范围

核对10份r1冻结输入哈希；分解26规格×7事件×2执行时点，共364条，逐日恒等式最大误差{v['daily_decomposition_error']:.2e}。随机日历2000次均为七段70日、年内段数一致；两种执行时点共18组直接重算验证向量统计，最大误差{v['random_max_direct_error']:.2e}。另核对EW对称收益等于多头腿加空头腿，最大误差{maxsplit:.2e}。

年度表零波动年份夏普保留为空，不赋予0。本轮不是独立样本外验证；不补齐r1未覆盖的8项研究、不新增宏观事件、不改生产。Excel和CSV均提供，随机日历与重点候选全部2000次结果可重查。方案见inputs/design.md。
'''
report=report.replace('双方法律/新闻事前知识不在数据里。','双方是否提前掌握或理解新闻，不在数据里。')
(O/'REPORT.md').write_text(report)
from openpyxl import load_workbook
wb=load_workbook(O/'事件依赖第二轮_候选复核.xlsx',read_only=True,data_only=True)
assert wb['复核结论'].max_row==27 and wb['事件方向与响应'].max_row==365
wb.close()
assert (O/'random_calendar_diagnostic.png').stat().st_size>10000
v['same_mapping_lf_review_count']=24;v['same_mapping_lf_neutral_leaders']=0;v['long_short_identity_max_error']=maxsplit;v['workbook_verified']=True
(O/'verification.json').write_text(json.dumps(v,indent=2))
shutil.copyfile('/tmp/event_dependence_r2.log',RUN/'analysis.log')
write_manifest(RUN,{'status':'complete','scope':'26 focused specifications; attribution and random-calendar diagnostic; r1 coverage gaps unchanged',
 'source_run':'20260915-event-dependence-r1','verification':v,'production_changed':False,
 'artifacts':[artifact_record(f,RUN) for f in sorted(RUN.rglob('*')) if f.is_file() and f.name!='manifest.json' and '__pycache__' not in f.parts]})
print('PACKAGED',json.dumps(v,indent=2))
print(countsummary.to_string(index=False))
